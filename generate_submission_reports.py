import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from reportlab.lib.pagesizes import letter
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Preformatted, PageBreak, KeepTogether
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

def set_cell_background(cell, fill_color):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_color)
    tcPr.append(shd)

def generate_docx(filename):
    doc = docx.Document()
    
    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
    # Styles
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = title_p.add_run("PES UNIVERSITY\nDEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING\n")
    r_title.bold = True
    r_title.font.size = Pt(14)
    r_title.font.color.rgb = RGBColor(16, 44, 87)
    
    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = sub_p.add_run("LAB ASSIGNMENT: DATA VERSION CONTROL (DVC) AND MLFLOW")
    r_sub.bold = True
    r_sub.font.size = Pt(16)
    r_sub.font.color.rgb = RGBColor(30, 86, 160)
    
    # Student metadata table
    tbl = doc.add_table(rows=4, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta = [
        ("Student Name:", "Aarav Yuval B G"),
        ("SRN:", "PES1UG24AM004"),
        ("Faculty / Evaluator:", "Sunaina Kumar"),
        ("Course / Semester:", "Software Engineering Lab (5th Semester)")
    ]
    for i, (k, v) in enumerate(meta):
        cell_k = tbl.cell(i, 0)
        cell_v = tbl.cell(i, 1)
        cell_k.text = k
        cell_v.text = v
        cell_k.paragraphs[0].runs[0].bold = True
        set_cell_background(cell_k, "F0F4F8")
        set_cell_background(cell_v, "FAFCFF")
        cell_k.width = Inches(2.2)
        cell_v.width = Inches(4.5)
        
    doc.add_paragraph().paragraph_format.space_after = Pt(12)
    
    def add_sec_heading(title):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(title)
        r.bold = True
        r.font.size = Pt(13)
        r.font.color.rgb = RGBColor(16, 44, 87)

    def add_sub_heading(title):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(title)
        r.bold = True
        r.font.size = Pt(11)
        r.font.color.rgb = RGBColor(40, 70, 110)

    # 1. Objectives
    add_sec_heading("1. Executive Summary & Lab Objectives")
    doc.add_paragraph(
        "Modern Machine Learning systems require tracking two fundamentally different assets: Code and Data. "
        "Standard Git version control is optimized for text source code but chokes on large datasets and binary weights. "
        "This laboratory assignment demonstrates the complete end-to-end MLOps pipeline using two synergistic tools:\n"
        "• Data Version Control (DVC): Versions large files, computes MD5 content hashes, isolates heavy data into content-addressable cache, and provides 'double-checkout' version switching.\n"
        "• MLflow: Manages experiment tracking, hyperparameter tuning, metric logging, model packaging, and inference serving.\n"
        "• The MLOps Bridge: Logging the exact DVC dataset hash as an MLflow parameter, establishing complete data lineage and reproducibility."
    )

    # 2. DVC Implementation
    add_sec_heading("2. Part 1: Data Version Control (DVC) Implementation")
    doc.add_paragraph(
        "A Git repository was initialized, followed by DVC initialization via 'dvc init'. "
        "The DVC workflow decouples dataset binaries from Git by creating lightweight pointer files (.dvc) while appending raw datasets to .gitignore."
    )
    
    add_sub_heading("2.1 Execution Steps & Commands")
    steps_table = doc.add_table(rows=6, cols=3)
    steps_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Stage", "Commands Executed", "Outcome / Recorded Value"]
    rows_data = [
        ("Init", "git init\ndvc init\ngit commit -m 'Initialize DVC'", "Created .dvc/ directory, .dvcignore, and Git initial commit."),
        ("Version 1.0", "echo 'version 1' > data.txt\ndvc add data.txt\ngit add data.txt.dvc .gitignore\ngit commit -m 'Add v1'\ngit tag -a 'v1.0'", "MD5 Hash: 70714828ec652824f0bd76e1e2da849e\nStored in .dvc/cache/files/md5/70/"),
        ("Version 2.0", "echo 'version 2' > data.txt\ndvc add data.txt\ngit add data.txt.dvc\ngit commit -m 'Update v2'\ngit tag -a 'v2.0'", "MD5 Hash: 75dafceb0d26a67df209575587f1334a\nPointer data.txt.dvc updated with new checksum."),
        ("Double Checkout", "git checkout v1.0\ndvc checkout\ncat data.txt", "Workspace data.txt restored to 'version 1' instantaneously from local DVC cache."),
        ("Remote Storage", "dvc remote add -d localremote <path>\ndvc push", "Synchronized data cache with local remote storage simulating cloud backup.")
    ]
    for c_idx, h in enumerate(headers):
        cell = steps_table.cell(0, c_idx)
        cell.text = h
        cell.paragraphs[0].runs[0].bold = True
        set_cell_background(cell, "E2ECF7")
        
    for r_idx, (st, cmd, out) in enumerate(rows_data):
        c0 = steps_table.cell(r_idx + 1, 0)
        c1 = steps_table.cell(r_idx + 1, 1)
        c2 = steps_table.cell(r_idx + 1, 2)
        c0.text = st
        c1.text = cmd
        c2.text = out
        c0.paragraphs[0].runs[0].bold = True
        set_cell_background(c0, "F8FAFC")
        set_cell_background(c1, "FFFFFF")
        set_cell_background(c2, "F8FAFC")
        c0.width = Inches(1.3)
        c1.width = Inches(2.7)
        c2.width = Inches(2.7)

    # 3. MLflow Implementation
    add_sec_heading("3. Part 2: MLflow Experiment Tracking & MLOps Integration")
    doc.add_paragraph(
        "MLflow was installed and configured to track machine learning experiments across multiple paradigms. "
        "The 4 core architectural components of MLflow were implemented:\n"
        "1. MLflow Tracking: Capturing parameters, metrics, tags, and run history.\n"
        "2. MLflow Projects: Standardized environment packaging.\n"
        "3. MLflow Models: Saving models in flavor formats (e.g. Scikit-Learn with cloudpickle serialization).\n"
        "4. MLflow Model Registry / Inference: Querying experiments dynamically by run ID and serving live predictions."
    )
    
    add_sub_heading("3.1 Implementation Modules Summary")
    doc.add_paragraph(
        "• 02_mlflow_dvc_pipeline.py: Bridge module reading the DVC hash (21d441a28bce4417276097df955afc50) of data/train.csv. "
        "Logged two runs: Run_1_Baseline (n_estimators=30, max_depth=2, Accuracy=0.9211) and Run_2_Tuned (n_estimators=120, max_depth=5, Accuracy=0.9211).\n"
        "• 03_mlflow_autolog.py: Demonstrated single-line autologging with mlflow.sklearn.autolog() for Logistic Regression.\n"
        "• 04_mlflow_wine.py: Explicit tracking on 13-feature Wine classification dataset using Random Forest, achieving 100% test accuracy.\n"
        "• 05_test_inference.py: Restores the trained model artifact from 'runs:/<run_id>/random_forest_model' and executes live inference."
    )

    # 4. Results Table
    add_sec_heading("4. Key Results & Metrics Summary")
    res_table = doc.add_table(rows=5, cols=4)
    res_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    r_headers = ["Experiment Name", "Run Name", "Logged Hyperparameters", "Recorded Accuracy"]
    r_rows = [
        ("DVC_MLflow_Pipeline", "Run_1_Baseline", "n_estimators=30, max_depth=2\nDVC Hash: 21d441a28bce...", "92.11%"),
        ("DVC_MLflow_Pipeline", "Run_2_Tuned", "n_estimators=120, max_depth=5\nDVC Hash: 21d441a28bce...", "92.11%"),
        ("MLflow_Quickstart", "Autolog_LogisticRegression", "solver='lbfgs', max_iter=1000", "100.0%"),
        ("Wine_Classification", "Wine_RandomForest_Run", "n_estimators=100, max_depth=3", "100.0%")
    ]
    for c_idx, h in enumerate(r_headers):
        cell = res_table.cell(0, c_idx)
        cell.text = h
        cell.paragraphs[0].runs[0].bold = True
        set_cell_background(cell, "E2ECF7")
        
    for r_idx, row in enumerate(r_rows):
        for c_idx, val in enumerate(row):
            cell = res_table.cell(r_idx + 1, c_idx)
            cell.text = val
            if c_idx == 0 or c_idx == 3:
                cell.paragraphs[0].runs[0].bold = True
            set_cell_background(cell, "FFFFFF" if r_idx % 2 == 0 else "F8FAFC")

    # 5. Conclusion
    add_sec_heading("5. Conclusion & Takeaways")
    doc.add_paragraph(
        "By integrating DVC with MLflow, complete MLOps reproducibility was achieved. "
        "Git manages lightweight metadata and source code, DVC versions and stores multi-version raw data in external storage, "
        "and MLflow monitors hyperparameter performance and model artifacts. "
        "The double-checkout mechanism provides instant data rollbacks, and the MLflow UI (port 5000) provides transparency for team collaboration."
    )
    
    doc.save(filename)
    print(f"Generated DOCX: {filename}")

def generate_pdf(filename):
    doc = SimpleDocTemplate(
        filename,
        pagesize=letter,
        rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40
    )
    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=15,
        leading=18,
        alignment=1,
        textColor=colors.HexColor('#102C57')
    )
    
    sub_title_style = ParagraphStyle(
        'SubTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=12,
        leading=15,
        alignment=1,
        textColor=colors.HexColor('#1E56A0')
    )
    
    h1_style = ParagraphStyle(
        'H1',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=11,
        leading=14,
        textColor=colors.HexColor('#102C57'),
        spaceBefore=10,
        spaceAfter=4
    )
    
    body_style = ParagraphStyle(
        'Body',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=12,
        textColor=colors.HexColor('#222222'),
        spaceAfter=6
    )
    
    code_style = ParagraphStyle(
        'CodeText',
        parent=styles['Normal'],
        fontName='Courier',
        fontSize=7.5,
        leading=9.5,
        textColor=colors.HexColor('#1A1A1A')
    )
    
    story = []
    story.append(Paragraph("PES UNIVERSITY", title_style))
    story.append(Paragraph("DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING", title_style))
    story.append(Spacer(1, 4))
    story.append(Paragraph("LAB ASSIGNMENT: DATA VERSION CONTROL (DVC) AND MLFLOW", sub_title_style))
    story.append(Spacer(1, 10))
    
    # Metadata Table
    meta_data = [
        [Paragraph("<b>Student Name:</b>", body_style), Paragraph("Aarav Yuval B G", body_style)],
        [Paragraph("<b>SRN:</b>", body_style), Paragraph("PES1UG24AM004", body_style)],
        [Paragraph("<b>Faculty / Evaluator:</b>", body_style), Paragraph("Sunaina Kumar", body_style)],
        [Paragraph("<b>Course / Semester:</b>", body_style), Paragraph("Software Engineering Lab (5th Semester)", body_style)]
    ]
    t_meta = Table(meta_data, colWidths=[140, 390])
    t_meta.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (0,-1), colors.HexColor('#EEF2F6')),
        ('BACKGROUND', (1,0), (1,-1), colors.HexColor('#F8FAFC')),
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#CBD5E1')),
        ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor('#E2E8F0')),
        ('TOPPADDING', (0,0), (-1,-1), 3),
        ('BOTTOMPADDING', (0,0), (-1,-1), 3),
    ]))
    story.append(t_meta)
    story.append(Spacer(1, 10))
    
    # 1. Executive Summary
    story.append(Paragraph("1. Executive Summary & Lab Objectives", h1_style))
    story.append(Paragraph(
        "Machine Learning operations (MLOps) necessitate versioning both source code and heavy data assets. "
        "Git fails when handling multi-megabyte/gigabyte binary datasets due to repository bloat and file limits. "
        "This laboratory assignment demonstrates the complementary use of <b>Data Version Control (DVC)</b> and <b>MLflow</b>. "
        "DVC provides content-addressable storage and lightweight Git pointers, while MLflow tracks experiments, "
        "hyperparameters, performance metrics, and packaged models. Linking the DVC dataset checksum into MLflow achieves 100% reproducible pipelines.",
        body_style
    ))
    
    # 2. DVC Implementation
    story.append(Paragraph("2. Part 1: Data Version Control (DVC) Implementation", h1_style))
    story.append(Paragraph(
        "DVC was initialized with <code>dvc init</code>. A data versioning lifecycle was simulated using <code>data.txt</code>, "
        "tracking version updates, Git commits, version tags (<code>v1.0</code> and <code>v2.0</code>), and double-checkout rollbacks.",
        body_style
    ))
    
    dvc_table_data = [
        [Paragraph("<b>Stage</b>", body_style), Paragraph("<b>Commands Executed</b>", body_style), Paragraph("<b>Outcome / Checksum</b>", body_style)],
        [Paragraph("<b>Init</b>", body_style), Paragraph("<code>git init<br/>dvc init<br/>git commit -m 'Init DVC'</code>", code_style), Paragraph("Initialized .dvc/, created .dvcignore.", body_style)],
        [Paragraph("<b>Version 1.0</b>", body_style), Paragraph("<code>echo 'version 1' > data.txt<br/>dvc add data.txt<br/>git commit -m 'v1' & tag v1.0</code>", code_style), Paragraph("MD5: 70714828ec652824f0bd76e1e2da849e<br/>Saved to .dvc/cache.", body_style)],
        [Paragraph("<b>Version 2.0</b>", body_style), Paragraph("<code>echo 'version 2' > data.txt<br/>dvc add data.txt<br/>git commit -m 'v2' & tag v2.0</code>", code_style), Paragraph("MD5: 75dafceb0d26a67df209575587f1334a<br/>Pointer file updated.", body_style)],
        [Paragraph("<b>Double Checkout</b>", body_style), Paragraph("<code>git checkout v1.0<br/>dvc checkout</code>", code_style), Paragraph("Workspace synced to 'version 1' data instantaneously from cache.", body_style)],
        [Paragraph("<b>Remote Storage</b>", body_style), Paragraph("<code>dvc remote add -d localremote &lt;path&gt;<br/>dvc push</code>", code_style), Paragraph("Cache pushed to external storage directory.", body_style)]
    ]
    t_dvc = Table(dvc_table_data, colWidths=[90, 230, 210])
    t_dvc.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#E2ECF7')),
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#CBD5E1')),
        ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor('#E2E8F0')),
        ('TOPPADDING', (0,0), (-1,-1), 3),
        ('BOTTOMPADDING', (0,0), (-1,-1), 3),
    ]))
    story.append(t_dvc)
    story.append(Spacer(1, 10))
    
    # 3. MLflow Implementation
    story.append(Paragraph("3. Part 2: MLflow Experiment Tracking & MLOps Bridge", h1_style))
    story.append(Paragraph(
        "Four modules were developed and executed to showcase the full MLflow capability suite:<br/>"
        "• <b>02_mlflow_dvc_pipeline.py:</b> Retrieves DVC hash (<code>21d441a28bce4417276097df955afc50</code>) and logs it into MLflow. Runs baseline vs tuned models.<br/>"
        "• <b>03_mlflow_autolog.py:</b> Demonstrates zero-effort tracking via <code>mlflow.sklearn.autolog()</code>.<br/>"
        "• <b>04_mlflow_wine.py:</b> Explicit logging on 13-feature Wine dataset achieving 100% test accuracy.<br/>"
        "• <b>05_test_inference.py:</b> Restores serialized model artifact from MLflow storage and executes live inference predictions.",
        body_style
    ))
    
    # 4. Results
    story.append(Paragraph("4. Key Results & Experiment Logs", h1_style))
    results_data = [
        [Paragraph("<b>Experiment</b>", body_style), Paragraph("<b>Run Name</b>", body_style), Paragraph("<b>Logged Params (inc. DVC Hash)</b>", body_style), Paragraph("<b>Accuracy</b>", body_style)],
        [Paragraph("<b>DVC_MLflow_Pipeline</b>", body_style), Paragraph("Run_1_Baseline", body_style), Paragraph("n_estimators=30, max_depth=2<br/>DVC: 21d441a28bce...", code_style), Paragraph("<b>92.11%</b>", body_style)],
        [Paragraph("<b>DVC_MLflow_Pipeline</b>", body_style), Paragraph("Run_2_Tuned", body_style), Paragraph("n_estimators=120, max_depth=5<br/>DVC: 21d441a28bce...", code_style), Paragraph("<b>92.11%</b>", body_style)],
        [Paragraph("<b>MLflow_Quickstart</b>", body_style), Paragraph("Autolog_LogisticRegression", body_style), Paragraph("solver='lbfgs', max_iter=1000", code_style), Paragraph("<b>100.0%</b>", body_style)],
        [Paragraph("<b>Wine_Classification</b>", body_style), Paragraph("Wine_RandomForest_Run", body_style), Paragraph("n_estimators=100, max_depth=3", code_style), Paragraph("<b>100.0%</b>", body_style)]
    ]
    t_res = Table(results_data, colWidths=[120, 110, 230, 70])
    t_res.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#E2ECF7')),
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#CBD5E1')),
        ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor('#E2E8F0')),
        ('TOPPADDING', (0,0), (-1,-1), 3),
        ('BOTTOMPADDING', (0,0), (-1,-1), 3),
    ]))
    story.append(t_res)
    story.append(Spacer(1, 10))
    
    # 5. Conclusion
    story.append(Paragraph("5. Summary & Viva Insights", h1_style))
    story.append(Paragraph(
        "<b>Key Takeaway:</b> DVC handles data versioning without Git bloat using MD5 checksums. "
        "MLflow handles experiment metrics and models. Together, logging the DVC hash into MLflow guarantees total reproducibility. "
        "The double checkout (<code>git checkout &lt;tag&gt; && dvc checkout</code>) seamlessly restores corresponding data states.",
        body_style
    ))
    
    doc.build(story)
    print(f"Generated PDF: {filename}")

if __name__ == "__main__":
    out_dir = "submissions"
    docx_path = os.path.join(out_dir, "PES1UG24AM004_DVC_and_MLflow_Report.docx")
    pdf_path = os.path.join(out_dir, "PES1UG24AM004_DVC_and_MLflow_Report.pdf")
    
    generate_docx(docx_path)
    generate_pdf(pdf_path)
    print("Reports generated successfully!")
